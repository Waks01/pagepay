"""Check that a Word document's prose, tables and pictures all come out.

The .docx path used to read the paragraph text and nothing else, so
every diagram in a Word document was dropped. This builds a document
with a heading, prose, a table and a picture, and checks that all
four survive.
"""
import io
import sys

sys.path.insert(0, ".")

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

from PIL import Image  # noqa: E402

from app.services import page_artifacts  # noqa: E402


def build_docx() -> bytes:
    import docx

    document = docx.Document()
    document.add_heading("Quadratic Equations", level=1)
    document.add_paragraph(
        "A quadratic equation is an equation of the form ax^2 + bx + c = 0, "
        "where a, b and c are constants and a is not zero."
    )
    document.add_paragraph("The roots of a quadratic equation are its zeros.")

    table = document.add_table(rows=2, cols=3)
    headers = ("Stage", "Input", "Output")
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    table.rows[1].cells[0].text = "Light reaction"
    table.rows[1].cells[1].text = "Light, H2O"
    table.rows[1].cells[2].text = "ATP, O2"

    # A picture, so the extraction has something to find.
    image = Image.new("RGB", (400, 200), "white")
    for x in range(400):
        for y in range(200):
            if abs((x - 200) ** 2 / 1000 - (y - 100)) < 3:
                image.putpixel((x, y), (20, 90, 160))
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    document.add_picture(io.BytesIO(buffer.getvalue()), width=3000000)
    document.add_paragraph("Figure 1: A parabola showing its vertex.")

    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def main() -> None:
    contents = build_docx()
    text, images = page_artifacts.extract_docx_text_and_images(contents)

    failures = []

    print(f"text: {len(text)} chars")
    print(f"images: {len(images)}")
    print()
    print(text)

    for expected in (
        "Quadratic Equations",
        "ax^2 + bx + c = 0",
        "[TABLE: Stage | Input | Output]",
        "Light reaction",
        "Figure 1: A parabola",
    ):
        if expected in text:
            print(f"ok   text has {expected!r}")
        else:
            failures.append(f"text is missing {expected!r}")
            print(f"FAIL text is missing {expected!r}")

    if len(images) == 1:
        print(f"ok   one picture came out ({len(images[0])} bytes)")
    else:
        failures.append(f"expected 1 picture, got {len(images)}")
        print(f"FAIL expected 1 picture, got {len(images)}")

    if images and images[0][:8] == b"\x89PNG\r\n\x1a\n":
        print("ok   the picture is a real PNG")
    else:
        failures.append("the picture is not a PNG")
        print("FAIL the picture is not a PNG")

    print()
    if failures:
        for failure in failures:
            print("FAILED:", failure)
        raise SystemExit(1)
    print("all docx checks passed")


if __name__ == "__main__":
    main()
