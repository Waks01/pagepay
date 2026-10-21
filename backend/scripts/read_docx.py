"""Read a .docx the way the study app reads one, and the way it could.

The app's current path uses python-docx and takes the paragraph text.
That drops every image in the document. This prints both: what the
app would see, and what is actually in the file.
"""
import io
import sys
import zipfile

import docx


PATH = r"..\words\MATH 102 - Algebra.docx"


def app_view(path: str) -> str:
    """Exactly what `_extract_document_text` produces for a .docx."""
    document = docx.Document(path)
    return "\n".join(para.text for para in document.paragraphs)


def main() -> None:
    path = PATH
    text = app_view(path)

    print("=" * 70)
    print(f"WHAT THE APP WOULD READ  ({len(text)} chars)")
    print("=" * 70)
    print(text[:2500])
    if len(text) > 2500:
        print(f"\n... [{len(text) - 2500} more characters] ...")
        print("\nTAIL:\n" + text[-1200:])

    print()
    print("=" * 70)
    print("WHAT IS ACTUALLY IN THE FILE")
    print("=" * 70)

    with zipfile.ZipFile(path) as archive:
        media = [n for n in archive.namelist() if n.startswith("word/media/")]
        print(f"embedded images: {len(media)}")
        for name in sorted(media):
            info = archive.getinfo(name)
            print(f"  {info.file_size:>8} bytes  {name}")

        document_xml = archive.read("word/document.xml").decode("utf-8", "replace")
        print(f"\ndocument.xml: {len(document_xml)} chars")
        print(f"  <w:drawing> elements : {document_xml.count('<w:drawing>')}")
        print(f"  <w:tbl> elements     : {document_xml.count('<w:tbl>')}")
        print(f"  <w:p> paragraphs     : {document_xml.count('<w:p ')}")

    # Equations: Word stores them as OMML, not as text.
    with zipfile.ZipFile(path) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8", "replace")
        print(f"  OMML equations       : {document_xml.count('<m:oMath')}")

    document = docx.Document(path)
    tables = document.tables
    print(f"\npython-docx tables: {len(tables)}")
    for index, table in enumerate(tables[:3]):
        print(f"  table {index}: {len(table.rows)} rows x {len(table.columns)} cols")
        for row in table.rows[:3]:
            print("    | " + " | ".join(cell.text.strip()[:40] for cell in row.cells))


if __name__ == "__main__":
    main()
